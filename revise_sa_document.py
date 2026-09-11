from pathlib import Path
import io, zipfile
from docx import Document
from PIL import Image, ImageDraw, ImageFont

SRC = Path(r"D:\SA\Doc\เอกสาร SA jingjing.docx")
OUT = Path(r"D:\SA\Mock-SA\เอกสาร SA jingjing ฉบับปรับให้ตรงระบบ.docx")
TMP = Path(r"D:\SA\Mock-SA\_revised_docx_tmp.docx")

def font(size=34, bold=False):
    candidates = [
        r"C:\Windows\Fonts\tahomabd.ttf" if bold else r"C:\Windows\Fonts\tahoma.ttf",
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
    ]
    for p in candidates:
        if Path(p).exists(): return ImageFont.truetype(p, size)
    return ImageFont.load_default()

def fit_text(draw, text, box, max_size=34, min_size=18, bold=False):
    x1,y1,x2,y2=box
    words=text.split(); lines=[]
    for size in range(max_size,min_size-1,-1):
        f=font(size,bold); lines=[]; cur=""
        for w in words:
            trial=(cur+" "+w).strip()
            if draw.textbbox((0,0),trial,font=f)[2] <= x2-x1-24: cur=trial
            else:
                if cur: lines.append(cur)
                cur=w
        if cur: lines.append(cur)
        lh=size+7
        if len(lines)*lh <= y2-y1-14: break
    total=len(lines)*lh
    for i,line in enumerate(lines):
        bb=draw.textbbox((0,0),line,font=f)
        draw.text(((x1+x2-bb[2])/2,y1+(y2-y1-total)/2+i*lh),line,font=f,fill="black")

def actor(draw, x, y, label, scale=1.0):
    r=int(22*scale); lw=max(3,int(5*scale))
    draw.ellipse((x-r,y-r,x+r,y+r),outline="black",width=lw)
    draw.line((x,y+r,x,y+100*scale),fill="black",width=lw)
    draw.line((x-45*scale,y+55*scale,x+45*scale,y+55*scale),fill="black",width=lw)
    draw.line((x,y+100*scale,x-40*scale,y+155*scale),fill="black",width=lw)
    draw.line((x,y+100*scale,x+40*scale,y+155*scale),fill="black",width=lw)
    bb=draw.textbbox((0,0),label,font=font(int(27*scale),True))
    draw.text((x-bb[2]/2,y+165*scale),label,font=font(int(27*scale),True),fill="black")

def usecase_diagram(size, title, actors, cases, relations):
    w,h=size; im=Image.new("RGB",size,"white"); d=ImageDraw.Draw(im)
    d.rounded_rectangle((int(w*.18),int(h*.04),int(w*.82),int(h*.96)),radius=30,outline="black",width=6)
    bb=d.textbbox((0,0),title,font=font(max(30,int(w*.018)),True))
    d.rectangle((w*.5-bb[2]/2-30,h*.025,w*.5+bb[2]/2+30,h*.075),fill="white",outline="black",width=3)
    d.text((w*.5-bb[2]/2,h*.04),title,font=font(max(30,int(w*.018)),True),fill="black")
    for label,x,y,links in actors:
        actor(d,int(x*w),int(y*h),label,max(.75,w/2700))
        for target in links:
            cx,cy,cw,ch=cases[target][1:]
            edge_x=(cx-cw/2)*w if x < .5 else (cx+cw/2)*w
            d.line((x*w+(70 if x<.5 else -70),y*h+70,edge_x,cy*h),fill="#555",width=4)
    for key,(label,cx,cy,cw,ch) in cases.items():
        box=((cx-cw/2)*w,(cy-ch/2)*h,(cx+cw/2)*w,(cy+ch/2)*h)
        d.ellipse(box,fill="white",outline="black",width=5)
        fit_text(d,label,box,max_size=max(25,int(w*.014)),min_size=17)
    for a,b,label in relations:
        _,ax,ay,aw,ah=cases[a]; _,bx,by,bw,bh=cases[b]
        d.line((ax*w,ay*h,bx*w,by*h),fill="#555",width=3)
        mx=(ax+bx)*w/2; my=(ay+by)*h/2
        d.text((mx,my),label,font=font(max(18,int(w*.009))),fill="#333")
    bio=io.BytesIO(); im.save(bio,"PNG"); return bio.getvalue()

def flow_diagram(size, title, lanes, nodes, arrows):
    w,h=size; im=Image.new("RGB",size,"white"); d=ImageDraw.Draw(im)
    margin=int(w*.04); header=int(h*.055); lane_w=(w-2*margin)/len(lanes)
    d.rectangle((margin,margin,w-margin,h-margin),outline="black",width=5)
    d.rectangle((margin,margin,w-margin,margin+header),fill="#e9eef5",outline="black",width=4)
    for i,l in enumerate(lanes):
        x=margin+i*lane_w
        if i: d.line((x,margin,x,h-margin),fill="#888",width=3)
        fit_text(d,l,(x,margin,x+lane_w,margin+header),max_size=max(24,int(w*.012)),bold=True)
    for key,(lane,y,label,kind) in nodes.items():
        cx=margin+(lane+.5)*lane_w; cy=y*h; bw=lane_w*.72; bh=h*.045
        if kind=="start":
            d.ellipse((cx-20,cy-20,cx+20,cy+20),fill="black")
        elif kind=="end":
            d.ellipse((cx-25,cy-25,cx+25,cy+25),outline="black",width=5); d.ellipse((cx-14,cy-14,cx+14,cy+14),fill="black")
        elif kind=="decision":
            pts=[(cx,cy-bh*.65),(cx+bw*.45,cy),(cx,cy+bh*.65),(cx-bw*.45,cy)]
            d.polygon(pts,fill="white",outline="black"); fit_text(d,label,(cx-bw*.35,cy-bh*.5,cx+bw*.35,cy+bh*.5),max_size=max(19,int(w*.009)))
        else:
            box=(cx-bw/2,cy-bh/2,cx+bw/2,cy+bh/2)
            d.rounded_rectangle(box,radius=18,fill="white",outline="black",width=4); fit_text(d,label,box,max_size=max(20,int(w*.009)))
    for a,b,label in arrows:
        la,ya,_,_=nodes[a]; lb,yb,_,_=nodes[b]
        x1=margin+(la+.5)*lane_w; x2=margin+(lb+.5)*lane_w; yy1=ya*h; yy2=yb*h
        d.line((x1,yy1+22,x2,yy2-22),fill="#333",width=4)
        if label: d.text(((x1+x2)/2+8,(yy1+yy2)/2),label,font=font(max(18,int(w*.008))),fill="#333")
    bio=io.BytesIO(); im.save(bio,"PNG"); return bio.getvalue()

def set_table(table, rows):
    for r, values in enumerate(rows):
        if r >= len(table.rows): break
        for c, value in enumerate(values):
            if c < len(table.rows[r].cells):
                table.rows[r].cells[c].text = value

d=Document(SRC)

# ระบบจัดการผู้ใช้งานและสิทธิ์
set_table(d.tables[1], [
 ["Stakeholder","บทบาทและหน้าที่","Internal External","Operational Executive"],
 ["เจ้าหน้าที่แต่ละฝ่าย","เข้าสู่ระบบ ตั้งรหัสผ่านครั้งแรก ดู/แก้ไขข้อมูลตนเอง เปลี่ยนรหัสผ่าน ส่งคำขอรีเซ็ตรหัสผ่าน และใช้โมดูลตามสิทธิ์","Internal","Operational"],
 ["ผู้ดูแลระบบ","สร้าง ค้นหา แก้ไข และปิดใช้งานบัญชีพนักงาน กำหนดสิทธิ์รายโมดูล พิจารณาคำขอรีเซ็ตรหัสผ่าน และตรวจประวัติกิจกรรม","Internal","Operational"],
])
set_table(d.tables[2], [
 ["ลำดับ","รายละเอียด Functional Requirement","Stakeholder ที่เกี่ยวข้อง"],
 ["1","ระบบต้องให้พนักงานเข้าสู่ระบบ ออกจากระบบ ตั้งรหัสผ่านครั้งแรก เปลี่ยนรหัสผ่าน และจัดการข้อมูลบัญชีของตนเอง","เจ้าหน้าที่แต่ละฝ่าย"],
 ["2","ระบบต้องให้ผู้ดูแลระบบสร้าง ค้นหา ดู แก้ไข และปิดใช้งานบัญชีพนักงาน","ผู้ดูแลระบบ"],
 ["3","ระบบต้องให้ผู้ดูแลระบบกำหนดสิทธิ์รายโมดูลเป็น none, view หรือ edit โดยใช้ค่าเริ่มต้นตามบทบาทและฝ่ายได้","ผู้ดูแลระบบ"],
 ["4","ระบบต้องรองรับคำขอรีเซ็ตรหัสผ่าน การพิจารณาโดยผู้ดูแลระบบ และการค้นดูประวัติกิจกรรมตามข้อมูลที่ระบบบันทึก","ผู้ดูแลระบบ, เจ้าหน้าที่แต่ละฝ่าย"],
])
set_table(d.tables[3], [
 ["รหัส","As a Actor","I want to Action","So that Goal and Value"],
 ["U01","เจ้าหน้าที่แต่ละฝ่าย","เข้าสู่ระบบและใช้เฉพาะโมดูลตามสิทธิ์","ปฏิบัติงานตามหน้าที่"],
 ["U02","เจ้าหน้าที่แต่ละฝ่าย","ดู/แก้ไขโปรไฟล์ เปลี่ยนรหัสผ่าน และส่งคำขอรีเซ็ตรหัสผ่าน","ดูแลบัญชีของตนเอง"],
 ["U03","ผู้ดูแลระบบ","สร้าง ค้นหา แก้ไข และปิดใช้งานบัญชีพนักงาน","ควบคุมบัญชีผู้ปฏิบัติงาน"],
 ["U04","ผู้ดูแลระบบ","กำหนดสิทธิ์ none/view/edit แยกตามโมดูล","จำกัดการเข้าถึงให้เหมาะกับหน้าที่"],
 ["U05","ผู้ดูแลระบบ","พิจารณาคำขอรีเซ็ตรหัสผ่านและตรวจประวัติกิจกรรม","ติดตามและช่วยเหลือการใช้งานย้อนหลัง"],
])
set_table(d.tables[5], [["Business Use Case","B01 จัดเตรียมและดูแลบัญชีพนักงาน"],["Business Actor","เจ้าหน้าที่แต่ละฝ่าย, ผู้ดูแลระบบ"],["ขั้นตอนการทำงาน Step by step","1. ผู้ดูแลระบบสร้างบัญชีพร้อมข้อมูลฝ่าย ตำแหน่ง และบทบาท 2. ระบบกำหนดสิทธิ์เริ่มต้นตามบทบาท/ฝ่าย 3. พนักงานตั้งรหัสผ่านครั้งแรกจากโทเคน 4. พนักงานเข้าสู่ระบบและจัดการข้อมูลตนเอง 5. ผู้ดูแลระบบแก้ไขหรือปิดใช้งานบัญชีเมื่อจำเป็น 6. หากลืมรหัสผ่าน พนักงานส่งคำขอและผู้ดูแลระบบพิจารณา"]])
set_table(d.tables[6], [["Business Use Case","B02 กำหนดและควบคุมสิทธิ์รายโมดูล"],["Business Actor","ผู้ดูแลระบบ, เจ้าหน้าที่แต่ละฝ่าย"],["ขั้นตอนการทำงาน Step by step","1. ผู้ดูแลระบบเลือกพนักงาน 2. ตรวจบทบาท ฝ่าย และสิทธิ์ปัจจุบัน 3. กำหนดสิทธิ์แต่ละโมดูลเป็น none, view หรือ edit 4. บันทึกสิทธิ์ 5. ระบบใช้สิทธิ์ดังกล่าวควบคุมเมนูและการดำเนินการ 6. Dashboard เปิดดูได้เสมอ ส่วน Audit และ Employee Admin จำกัดเฉพาะผู้ดูแลระบบ"]])
set_table(d.tables[7], [["Business Use Case","B03 ตรวจสอบประวัติและคำขอรีเซ็ตรหัสผ่าน"],["Business Actor","ผู้ดูแลระบบ"],["ขั้นตอนการทำงาน Step by step","1. ผู้ดูแลระบบเปิดประวัติกิจกรรมหรือรายการคำขอรีเซ็ตรหัสผ่าน 2. ค้นหาหรือกรองรายการ 3. ระบบแสดงผู้ดำเนินการ การกระทำ โมดูล เป้าหมาย คำอธิบาย และวันเวลาเท่าที่บันทึกไว้ 4. ผู้ดูแลระบบพิจารณาคำขอรีเซ็ตรหัสผ่าน 5. ระบบบันทึกผลการพิจารณา"]])
set_table(d.tables[8], [
 ["รหัส","ชื่อ System Use Case","คำอธิบายโดยย่อ"],
 ["U01","พนักงานเข้าสู่ระบบและออกจากระบบ","ระบบตรวจบัญชี รหัสผ่าน สถานะ และโหลดสิทธิ์รายโมดูล"],
 ["U02","พนักงานจัดการบัญชีและรหัสผ่านของตนเอง","ดู/แก้ไขโปรไฟล์ เปลี่ยนรหัสผ่าน ตั้งรหัสผ่านครั้งแรก หรือส่งคำขอรีเซ็ต"],
 ["U03","ผู้ดูแลระบบจัดการบัญชีพนักงาน","สร้าง ค้นหา ดู แก้ไข และปิดใช้งานบัญชีพนักงาน"],
 ["U04","ผู้ดูแลระบบกำหนดสิทธิ์รายโมดูล","กำหนดระดับ none, view หรือ edit และใช้ค่าเริ่มต้นตามบทบาท/ฝ่าย"],
 ["U05","ผู้ดูแลระบบตรวจประวัติและพิจารณาคำขอรีเซ็ต","ค้นดูประวัติกิจกรรมและอนุมัติ/ปฏิเสธคำขอรีเซ็ตรหัสผ่าน"],
])
set_table(d.tables[9], [["System Use Case","U01 พนักงานเข้าสู่ระบบและออกจากระบบ"],["Actor","เจ้าหน้าที่แต่ละฝ่าย, ผู้ดูแลระบบ"],["Preconditions","1. มีบัญชีพนักงาน 2. บัญชีอยู่ในสถานะใช้งาน 3. หากเป็นบัญชีใหม่ต้องตั้งรหัสผ่านครั้งแรกแล้ว"],["ขั้นตอนการทำงาน Step by step","1. เปิดหน้าเข้าสู่ระบบ 2. กรอกอีเมล/ชื่อผู้ใช้และรหัสผ่าน 3. ระบบตรวจข้อมูลและสถานะบัญชี 4. ระบบสร้างเซสชันและโหลดสิทธิ์รายโมดูล 5. ระบบแสดงเมนูตามสิทธิ์ 6. เมื่อออกจากระบบ ระบบสิ้นสุดเซสชัน"],["Postconditions","เข้าสู่ระบบตามสิทธิ์ หรือออกจากระบบสำเร็จ และเกิดรายการกิจกรรมตามที่ระบบบันทึก"],["Abnormal Paths","ข้อมูลไม่ครบ รหัสผ่านผิด บัญชีไม่ใช้งาน หรือระบบตรวจสอบข้อมูลไม่ได้ ระบบปฏิเสธและแจ้งข้อผิดพลาด"]])
set_table(d.tables[10], [["System Use Case","U02 พนักงานจัดการบัญชีและรหัสผ่านของตนเอง"],["Actor","เจ้าหน้าที่แต่ละฝ่าย, ผู้ดูแลระบบ"],["Preconditions","ผู้ใช้มีบัญชี; การดู/แก้ไขโปรไฟล์และเปลี่ยนรหัสผ่านต้องเข้าสู่ระบบ ส่วนการตั้งรหัสผ่านครั้งแรกใช้โทเคน"],["ขั้นตอนการทำงาน Step by step","1. เปิดหน้าบัญชีของฉัน 2. ดูหรือแก้ไขข้อมูลที่ระบบอนุญาต 3. เปลี่ยนรหัสผ่านโดยยืนยันรหัสเดิม หรือใช้โทเคนตั้งรหัสผ่านครั้งแรก 4. หากลืมรหัสผ่านให้ส่งคำขอรีเซ็ต 5. ระบบตรวจและบันทึกผล"],["Postconditions","ข้อมูลส่วนตัวหรือรหัสผ่านถูกปรับปรุง หรือมีคำขอรีเซ็ตรหัสผ่านรอพิจารณา"],["Abnormal Paths","โทเคนไม่ถูกต้อง/หมดอายุ รหัสเดิมผิด ข้อมูลไม่ครบ หรือบันทึกไม่สำเร็จ ระบบไม่เปลี่ยนข้อมูล"]])
set_table(d.tables[11], [["System Use Case","U03 ผู้ดูแลระบบจัดการบัญชีพนักงาน"],["Actor","ผู้ดูแลระบบ"],["Preconditions","ผู้ดูแลระบบเข้าสู่ระบบแล้ว"],["ขั้นตอนการทำงาน Step by step","1. เปิดหน้าจัดการพนักงาน 2. ค้นหาหรือเลือกเพิ่มพนักงาน 3. กรอก/แก้ไขข้อมูล ชื่อ อีเมล เบอร์โทร ฝ่าย ตำแหน่ง และบทบาท 4. ระบบตรวจข้อมูลซ้ำและความถูกต้อง 5. บันทึกบัญชี หรือเลือกปิดใช้งานบัญชีเดิม 6. ระบบบันทึกกิจกรรม"],["Postconditions","บัญชีพนักงานถูกสร้าง แก้ไข หรือเปลี่ยนเป็นไม่ใช้งาน"],["Abnormal Paths","อีเมลซ้ำ ข้อมูลไม่ครบ ไม่พบบัญชี หรือบันทึกไม่สำเร็จ ระบบแจ้งข้อผิดพลาด"]])
set_table(d.tables[12], [["System Use Case","U04 ผู้ดูแลระบบกำหนดสิทธิ์รายโมดูล"],["Actor","ผู้ดูแลระบบ"],["Preconditions","ผู้ดูแลระบบเข้าสู่ระบบและบัญชีพนักงานเป้าหมายมีอยู่"],["ขั้นตอนการทำงาน Step by step","1. เปิดฟอร์มพนักงาน 2. เลือกบทบาทและฝ่ายเพื่อโหลดค่าเริ่มต้นได้ 3. ปรับแต่ละโมดูลเป็น none, view หรือ edit 4. บันทึก 5. ระบบใช้สิทธิ์ใหม่ในการควบคุมการเข้าถึง; Dashboard ดูได้เสมอ และ Audit/Employee Admin เป็นสิทธิ์ผู้ดูแลระบบ"],["Postconditions","สิทธิ์รายโมดูลของพนักงานถูกบันทึกและมีผลกับการใช้งาน"],["Abnormal Paths","ไม่พบบัญชีหรือบันทึกไม่สำเร็จ ระบบคงค่าเดิมและแจ้งข้อผิดพลาด"]])
set_table(d.tables[13], [["System Use Case","U05 ผู้ดูแลระบบตรวจประวัติและพิจารณาคำขอรีเซ็ต"],["Actor","ผู้ดูแลระบบ"],["Preconditions","ผู้ดูแลระบบเข้าสู่ระบบแล้ว"],["ขั้นตอนการทำงาน Step by step","1. เปิดหน้าประวัติกิจกรรมหรือรายการคำขอรีเซ็ต 2. ค้นหา/กรองรายการ 3. ระบบแสดง actor, action, module, target, description และวันเวลาเท่าที่มี 4. สำหรับคำขอรีเซ็ต ผู้ดูแลระบบเลือกอนุมัติหรือปฏิเสธ 5. ระบบบันทึกผล"],["Postconditions","ตรวจสอบกิจกรรมย้อนหลังได้ และคำขอรีเซ็ตมีผลการพิจารณา"],["Abnormal Paths","ไม่พบรายการหรือดึงข้อมูลไม่ได้ ระบบแสดงสถานะที่เหมาะสมโดยไม่เปลี่ยนข้อมูลอื่น"]])
set_table(d.tables[14], [["Actor","ประเภท","System Use Case ที่เกี่ยวข้อง"],["เจ้าหน้าที่แต่ละฝ่าย","Person","U01, U02"],["ผู้ดูแลระบบ","Person","U01, U02, U03, U04, U05"]])
set_table(d.tables[15], [["Use Case ต้นทาง","ความสัมพันธ์","Use Case ปลายทาง","เหตุผล"],["U03 จัดการบัญชีพนักงาน","include","U04 กำหนดสิทธิ์รายโมดูล","การสร้างหรือแก้ไขพนักงานสามารถกำหนดสิทธิ์รายโมดูลในฟอร์มเดียวกัน"]])

# ระบบจัดการโปรโมชั่น
set_table(d.tables[16], [["Stakeholder","บทบาทและหน้าที่","Internal External","Operational Executive"],["เจ้าหน้าที่ฝ่ายการตลาด","สร้าง ค้นหา ดู แก้ไข และลบแบบคงประวัติโปรโมชั่น ซึ่งการบันทึกจะสร้างคำขออนุมัติทันที","Internal","Operational"],["ผู้อนุมัติโปรโมชั่น","ตรวจรายการรออนุมัติและเลือกอนุมัติหรือปฏิเสธพร้อมหมายเหตุ","Internal","Executive"],["ลูกค้า","ดูโปรโมชั่นที่ใช้ได้ แล้วเลือกโปรโมชั่นหนึ่งรายการหรือกรอกรหัสในหน้าจองบัตร","External","Operational"]])
set_table(d.tables[17], [
 ["ลำดับ","รายละเอียด Functional Requirement","Stakeholder ที่เกี่ยวข้อง"],
 ["1","สร้าง ค้นหา ดู แก้ไข และลบแบบ soft delete โปรโมชั่น โดยผูกกับคอนเสิร์ตหนึ่งงานและโซนที่เลือก","เจ้าหน้าที่ฝ่ายการตลาด"],
 ["2","กำหนดชื่อ คำอธิบาย/เงื่อนไข แบนเนอร์ รหัสโปรโมชั่น ส่วนลดแบบเปอร์เซ็นต์หรือจำนวนเงิน ส่วนลดสูงสุด ยอดขั้นต่ำ ช่วงเวลา โควตารวม และจำนวนครั้งต่อคน","เจ้าหน้าที่ฝ่ายการตลาด"],
 ["3","ตรวจฟิลด์บังคับ รูปแบบ ช่วงเวลา การอ้างอิงคอนเสิร์ต/โซน และรหัสโปรโมชั่นไม่ซ้ำ","เจ้าหน้าที่ฝ่ายการตลาด"],
 ["4","เมื่อบันทึก ระบบสร้าง/ปรับโปรโมชั่นและคำขออนุมัติสถานะรออนุมัติทันที","เจ้าหน้าที่ฝ่ายการตลาด"],
 ["5","ผู้อนุมัติอนุมัติหรือปฏิเสธรายการรออนุมัติ พร้อมบันทึกหมายเหตุ ผู้พิจารณา และเวลา","ผู้อนุมัติโปรโมชั่น"],
 ["6","ลบโปรโมชั่นแบบ soft delete เพื่อไม่แสดงกับรายการใหม่ แต่ยังคงข้อมูลในฐานข้อมูล","เจ้าหน้าที่ฝ่ายการตลาด"],
 ["7","แสดงโปรโมชั่นที่ active และผ่านเงื่อนไขคอนเสิร์ต โซน ยอดขั้นต่ำ ช่วงเวลา และค่าโควตาที่มีอยู่","ลูกค้า"],
 ["8","ระบบเลือกโปรโมชั่นที่ผ่านเงื่อนไขรายการแรกเป็นค่าเริ่มต้น ลูกค้าเปลี่ยนรายการหรือกรอกรหัสได้ และใช้ได้ครั้งละหนึ่งโปรโมชั่น","ลูกค้า"],
 ["9","ระบบคำนวณส่วนลดจากโปรโมชั่นที่เลือกและส่งยอดส่วนลดไปกับคำขอสร้างการจอง","ลูกค้า"],
 ["10","เวอร์ชันปัจจุบันยังไม่เชื่อมการจองกับ PromotionUsageLog และยังไม่เพิ่ม/คืน UsedQuota อัตโนมัติ","ผู้เกี่ยวข้อง"],
 ["11","หน้าจัดการแสดงรายละเอียด สถานะ และประวัติคำขออนุมัติ; ไม่ได้สรุปยอดใช้จริงจากการจองในเวอร์ชันปัจจุบัน","เจ้าหน้าที่ฝ่ายการตลาด, ผู้อนุมัติโปรโมชั่น"],
])
set_table(d.tables[18], [
 ["รหัส","As a Actor","I want to Action","So that Goal and Value"],
 ["US01","เจ้าหน้าที่ฝ่ายการตลาด","สร้าง ดู แก้ไข และลบโปรโมชั่นตามฟิลด์ที่ระบบรองรับ","จัดการข้อมูลโปรโมชั่น"],
 ["US02","เจ้าหน้าที่ฝ่ายการตลาด","ให้ระบบตรวจรูปแบบ การอ้างอิง และรหัสไม่ซ้ำ","ลดข้อมูลผิดพลาดก่อนบันทึก"],
 ["US03","เจ้าหน้าที่ฝ่ายการตลาด","ให้การบันทึกสร้างคำขออนุมัติทันที","ส่งรายการเข้าสู่ขั้นตอนพิจารณา"],
 ["US04","ผู้อนุมัติโปรโมชั่น","อนุมัติหรือปฏิเสธพร้อมหมายเหตุ","ควบคุมโปรโมชั่นก่อนเปิดใช้งาน"],
 ["US05","เจ้าหน้าที่ฝ่ายการตลาด","ลบโปรโมชั่นแบบคงประวัติ","หยุดแสดงโปรโมชั่นกับรายการใหม่"],
 ["US06","ลูกค้า","เห็นโปรโมชั่นที่ผ่านเงื่อนไขของคอนเสิร์ต โซน และยอดซื้อ","เลือกส่วนลดที่ใช้ได้"],
 ["US07","ลูกค้า","เลือกโปรโมชั่นหรือกรอกรหัสโปรโมชั่นได้หนึ่งรายการ","ใช้โปรโมชั่นที่ต้องการกับการจอง"],
 ["US08","ลูกค้า","เห็นส่วนลดและยอดสุทธิที่คำนวณจากโปรโมชั่นที่เลือก","ตรวจสอบราคาก่อนจอง"],
 ["US09","เจ้าหน้าที่ฝ่ายการตลาดและผู้อนุมัติ","ดูสถานะและประวัติคำขออนุมัติ","ติดตามผลการพิจารณา"],
 ["US10","ผู้เกี่ยวข้อง","ทราบข้อจำกัดการบันทึกการใช้และโควตาในเวอร์ชันปัจจุบัน","ไม่อ้างความสามารถที่ยังไม่ได้เชื่อมกับการจอง"],
])
set_table(d.tables[19], [["รหัส","ชื่อ Business Use Case","Business Actor ที่เกี่ยวข้อง"],["B01","จัดทำและบันทึกโปรโมชั่น","เจ้าหน้าที่ฝ่ายการตลาด"],["B02","พิจารณาอนุมัติโปรโมชั่น","เจ้าหน้าที่ฝ่ายการตลาด, ผู้อนุมัติโปรโมชั่น"],["B03","เลือกหรือกรอกรหัสโปรโมชั่นระหว่างจองบัตร","ลูกค้าผู้ซื้อบัตร"],["B04","คำนวณส่วนลดในรายการจอง","ลูกค้าผู้ซื้อบัตร"],["B05","ติดตามสถานะและประวัติการอนุมัติ","เจ้าหน้าที่ฝ่ายการตลาด, ผู้อนุมัติโปรโมชั่น"]])
set_table(d.tables[20], [["Business Use Case","B01 จัดทำและบันทึกโปรโมชั่น"],["Business Actor","เจ้าหน้าที่ฝ่ายการตลาด"],["ขั้นตอนการทำงาน Step by step","1. เลือกคอนเสิร์ตหนึ่งงานและโซนที่ร่วมรายการ 2. กรอกชื่อ คำอธิบาย/เงื่อนไข แบนเนอร์ และรหัส 3. กำหนดประเภท/มูลค่าส่วนลด ส่วนลดสูงสุด และยอดขั้นต่ำ 4. กำหนดช่วงเวลา โควตารวม และจำนวนครั้งต่อคน 5. ระบบตรวจฟิลด์ รูปแบบ การอ้างอิง และรหัสไม่ซ้ำ 6. เมื่อบันทึก ระบบสร้างหรือปรับโปรโมชั่นพร้อมคำขออนุมัติสถานะรออนุมัติทันที"]])
set_table(d.tables[21], [["Business Use Case","B02 พิจารณาอนุมัติโปรโมชั่น"],["Business Actor","เจ้าหน้าที่ฝ่ายการตลาด, ผู้อนุมัติโปรโมชั่น"],["ขั้นตอนการทำงาน Step by step","1. ผู้อนุมัติเปิดรายการรออนุมัติ 2. ตรวจรายละเอียดโปรโมชั่น 3. เลือกอนุมัติหรือปฏิเสธและบันทึกหมายเหตุ 4. ระบบบันทึกผล ผู้พิจารณา และเวลา 5. โปรโมชั่นที่อนุมัติและ active สามารถแสดงตามช่วงเวลา 6. รายการที่ถูกปฏิเสธสามารถแก้ไขและบันทึกเพื่อสร้างคำขอรอบใหม่"]])
set_table(d.tables[22], [["Business Use Case","B03 เลือกหรือกรอกรหัสโปรโมชั่นระหว่างจองบัตร"],["Business Actor","ลูกค้าผู้ซื้อบัตร"],["ขั้นตอนการทำงาน Step by step","1. ลูกค้าเลือกคอนเสิร์ต โซน ที่นั่ง และเข้าสู่หน้าสรุป 2. ระบบโหลดโปรโมชั่นที่ active และผ่านเงื่อนไขช่วงเวลา โควตา คอนเสิร์ต โซน และยอดขั้นต่ำ 3. ระบบเลือกโปรโมชั่นที่ผ่านเงื่อนไขรายการแรกเป็นค่าเริ่มต้น 4. ลูกค้าเลือกโปรโมชั่นอื่นหรือกรอกรหัส 5. ระบบตรวจรหัสและเงื่อนไขอีกครั้ง 6. ใช้ได้ครั้งละหนึ่งโปรโมชั่น 7. ระบบแสดงส่วนลดและยอดสุทธิ 8. ลูกค้ายืนยันการจอง"]])
set_table(d.tables[23], [["Business Use Case","B04 คำนวณส่วนลดในรายการจอง"],["Business Actor","ลูกค้าผู้ซื้อบัตร"],["ขั้นตอนการทำงาน Step by step","1. รับโปรโมชั่นที่ลูกค้าเลือก 2. คำนวณส่วนลดแบบเปอร์เซ็นต์หรือจำนวนเงิน 3. จำกัดด้วยส่วนลดสูงสุดเมื่อกำหนด 4. แสดงยอดก่อนลด ส่วนลด และยอดสุทธิ 5. ส่งค่า discount_amount ไปกับคำขอสร้างการจอง 6. เวอร์ชันปัจจุบันไม่ได้ส่ง promotion_id และยังไม่บันทึก PromotionUsageLog หรือปรับ UsedQuota จากขั้นตอนจอง"]])
set_table(d.tables[24], [["Business Use Case","B05 ติดตามสถานะและประวัติการอนุมัติ"],["Business Actor","เจ้าหน้าที่ฝ่ายการตลาด, ผู้อนุมัติโปรโมชั่น"],["ขั้นตอนการทำงาน Step by step","1. เปิดรายการโปรโมชั่นหรือคำขออนุมัติ 2. ค้นหาและเลือกโปรโมชั่น 3. ดูรายละเอียดและสถานะ 4. ดูประวัติคำขออนุมัติ ผู้พิจารณา หมายเหตุ และเวลา 5. ใช้ข้อมูลติดตามการพิจารณา 6. ไม่สรุปยอดใช้จริงหรือการคืนโควตาจากการจอง เพราะเวอร์ชันปัจจุบันยังไม่เชื่อมข้อมูลดังกล่าว"]])
set_table(d.tables[25], [
 ["รหัส","ชื่อ System Use Case","คำอธิบายโดยย่อ"],
 ["U01","ฝ่ายการตลาดจัดการข้อมูลโปรโมชั่น","สร้าง ค้นหา ดู แก้ไข และลบแบบ soft delete"],
 ["U02","ระบบตรวจข้อมูลโปรโมชั่น","ตรวจฟิลด์ รูปแบบ ช่วงเวลา การอ้างอิง และรหัสไม่ซ้ำ"],
 ["U03","ระบบสร้างคำขออนุมัติเมื่อบันทึก","สร้างสถานะรออนุมัติทันทีเมื่อสร้างหรือแก้ไข"],
 ["U04","ผู้อนุมัติพิจารณาโปรโมชั่น","อนุมัติหรือปฏิเสธพร้อมหมายเหตุ"],
 ["U05","ฝ่ายการตลาดลบโปรโมชั่น","ลบแบบ soft delete และหยุดแสดงกับรายการใหม่"],
 ["U06","ระบบตรวจโปรโมชั่นที่ใช้ได้","กรองตาม active ช่วงเวลา โควตา คอนเสิร์ต โซน และยอดขั้นต่ำ"],
 ["U07","ลูกค้าเลือกหรือกรอกรหัสโปรโมชั่น","เลือกได้หนึ่งรายการ โดยรายการแรกที่ผ่านเงื่อนไขเป็นค่าเริ่มต้น"],
 ["U08","ระบบคำนวณส่วนลดสำหรับการจอง","คำนวณส่วนลดและส่ง discount_amount ไปกับการสร้างการจอง"],
 ["U09","ผู้เกี่ยวข้องตรวจสถานะและประวัติอนุมัติ","ดูรายละเอียด สถานะ และข้อมูลการพิจารณาที่ระบบบันทึก"],
])
set_table(d.tables[26], [["System Use Case","U01 ฝ่ายการตลาดจัดการข้อมูลโปรโมชั่น"],["Actor","เจ้าหน้าที่ฝ่ายการตลาด"],["Preconditions","เข้าสู่ระบบและมีสิทธิ์จัดการโปรโมชั่น"],["ขั้นตอนการทำงาน Step by step","1. เปิดรายการโปรโมชั่น 2. ค้นหาหรือเลือกสร้าง 3. กรอก/แก้ไขฟิลด์ที่รองรับ 4. เลือกคอนเสิร์ตหนึ่งงานและโซน 5. ระบบตรวจข้อมูล 6. บันทึกแล้วสร้างคำขออนุมัติทันที 7. หากลบ ระบบทำ soft delete"],["Postconditions","โปรโมชั่นถูกสร้าง แก้ไข หรือลบแบบคงข้อมูล และการบันทึกมีคำขอรออนุมัติ"],["Abnormal Paths","ข้อมูลไม่ครบ รูปแบบผิด อ้างอิงไม่ถูกต้อง รหัสซ้ำ หรือบันทึกไม่ได้ ระบบแจ้งข้อผิดพลาด"]])
set_table(d.tables[27], [["System Use Case","U02 ระบบตรวจข้อมูลโปรโมชั่น"],["Actor","เจ้าหน้าที่ฝ่ายการตลาด"],["Preconditions","มีข้อมูลโปรโมชั่นที่กำลังสร้างหรือแก้ไข"],["ขั้นตอนการทำงาน Step by step","1. ตรวจฟิลด์บังคับ 2. ตรวจชนิดและช่วงค่าตัวเลข 3. ตรวจวันเริ่ม/สิ้นสุด 4. ตรวจคอนเสิร์ตและโซน 5. ตรวจรหัสโปรโมชั่นไม่ซ้ำ 6. แจ้งรายการผิดพลาด หรืออนุญาตให้บันทึก"],["Postconditions","ข้อมูลที่ผ่านการตรวจพร้อมบันทึก"],["Abnormal Paths","ข้อมูลไม่ผ่าน ระบบไม่บันทึกและแสดงข้อผิดพลาด"]])
set_table(d.tables[28], [["System Use Case","U03 ระบบสร้างคำขออนุมัติเมื่อบันทึก"],["Actor","เจ้าหน้าที่ฝ่ายการตลาด"],["Preconditions","ข้อมูลโปรโมชั่นผ่านการตรวจ"],["ขั้นตอนการทำงาน Step by step","1. เจ้าหน้าที่กดบันทึก 2. ระบบสร้างหรือปรับข้อมูลโปรโมชั่น 3. ระบบสร้างคำขออนุมัติสถานะ pending 4. ระบบผูกคำขอกับโปรโมชั่นและบันทึกเวลา 5. แสดงผลสำเร็จ"],["Postconditions","มีคำขออนุมัติรอพิจารณาโดยอัตโนมัติ"],["Abnormal Paths","หากบันทึกส่วนใดไม่สำเร็จ ระบบคืนข้อผิดพลาดและไม่ถือว่าส่งคำขอสำเร็จ"]])
set_table(d.tables[29], [["System Use Case","U04 ผู้อนุมัติพิจารณาโปรโมชั่น"],["Actor","ผู้อนุมัติโปรโมชั่น"],["Preconditions","เข้าสู่ระบบและมีรายการสถานะ pending"],["ขั้นตอนการทำงาน Step by step","1. เปิดรายการรออนุมัติ 2. เลือกรายการ 3. ตรวจรายละเอียด 4. เลือกอนุมัติหรือปฏิเสธ 5. กรอกหมายเหตุ 6. ระบบบันทึกผล ผู้พิจารณา และเวลา 7. ปรับสถานะโปรโมชั่นตามผล"],["Postconditions","คำขอและโปรโมชั่นมีสถานะตามผลพิจารณา"],["Abnormal Paths","รายการไม่อยู่ในสถานะ pending หรือบันทึกไม่ได้ ระบบไม่เปลี่ยนผล"]])
set_table(d.tables[30], [["System Use Case","U05 ฝ่ายการตลาดลบโปรโมชั่น"],["Actor","เจ้าหน้าที่ฝ่ายการตลาด"],["Preconditions","เข้าสู่ระบบ มีสิทธิ์ และพบโปรโมชั่น"],["ขั้นตอนการทำงาน Step by step","1. ค้นหาโปรโมชั่น 2. เลือกลบ 3. ยืนยัน 4. ระบบทำ soft delete 5. รายการไม่ปรากฏในคำค้นปกติและไม่แสดงกับลูกค้า"],["Postconditions","โปรโมชั่นถูกลบแบบคงข้อมูลในฐานข้อมูล"],["Abnormal Paths","ไม่พบรายการหรือดำเนินการไม่ได้ ระบบแจ้งข้อผิดพลาด"]])
set_table(d.tables[31], [["System Use Case","U06 ระบบตรวจโปรโมชั่นที่ใช้ได้"],["Actor","ลูกค้าผู้ซื้อบัตร"],["Preconditions","ลูกค้าอยู่ในขั้นตอนเลือกที่นั่ง/สรุปรายการ และมีข้อมูลคอนเสิร์ต โซน และยอดรวม"],["ขั้นตอนการทำงาน Step by step","1. โหลดโปรโมชั่น active 2. ตรวจช่วงเวลา 3. ตรวจค่าโควตาที่ระบบมี 4. ตรวจคอนเสิร์ตและโซน 5. ตรวจยอดขั้นต่ำ 6. แสดงรายการที่ผ่านเงื่อนไข"],["Postconditions","ได้รายการโปรโมชั่นที่สามารถเลือกได้ในขณะนั้น"],["Abnormal Paths","หากไม่มีรายการ ระบบให้จองต่อโดยไม่ใช้โปรโมชั่น; หากโหลดไม่ได้ ระบบแจ้งข้อผิดพลาด"]])
set_table(d.tables[32], [["System Use Case","U07 ลูกค้าเลือกหรือกรอกรหัสโปรโมชั่น"],["Actor","ลูกค้าผู้ซื้อบัตร"],["Preconditions","อยู่ในหน้าสรุปรายการจอง"],["ขั้นตอนการทำงาน Step by step","1. ระบบเลือกโปรโมชั่นที่ผ่านเงื่อนไขรายการแรกเป็นค่าเริ่มต้นถ้ามี 2. ลูกค้าเลือกโปรโมชั่นอื่นจากรายการ หรือกรอกรหัส 3. ระบบตรวจรหัสและเงื่อนไข 4. ใช้ได้ครั้งละหนึ่งโปรโมชั่น 5. คำนวณและแสดงส่วนลดกับยอดสุทธิ 6. ลูกค้ายืนยันการจอง"],["Postconditions","รายการจองมีส่วนลดจากโปรโมชั่นที่เลือกหนึ่งรายการ หรือไม่มีโปรโมชั่น"],["Abnormal Paths","รหัสไม่ถูกต้อง ไม่ผ่านเงื่อนไข หรือโปรโมชั่นใช้ไม่ได้ ระบบแจ้งและไม่ใช้ส่วนลดนั้น"]])
set_table(d.tables[33], [["System Use Case","U08 ระบบคำนวณส่วนลดสำหรับการจอง"],["Actor","ลูกค้าผู้ซื้อบัตร"],["Preconditions","มีโปรโมชั่นที่เลือกและผ่านเงื่อนไข"],["ขั้นตอนการทำงาน Step by step","1. รับยอดรวมและข้อมูลส่วนลด 2. คำนวณแบบเปอร์เซ็นต์หรือจำนวนเงิน 3. ใช้เพดานส่วนลดสูงสุดเมื่อกำหนด 4. แสดงยอดสุทธิ 5. ส่ง discount_amount ไปยังการสร้างการจอง 6. ไม่บันทึก promotion_id/PromotionUsageLog และไม่ปรับ UsedQuota ในเวอร์ชันปัจจุบัน"],["Postconditions","การจองได้รับค่า discount_amount ที่คำนวณแล้ว"],["Abnormal Paths","คำนวณหรือสร้างการจองไม่ได้ ระบบแจ้งข้อผิดพลาด; ไม่มีการอ้างว่าตัดหรือคืนโควตาอัตโนมัติ"]])
set_table(d.tables[34], [["System Use Case","U09 ผู้เกี่ยวข้องตรวจสถานะและประวัติอนุมัติ"],["Actor","เจ้าหน้าที่ฝ่ายการตลาด, ผู้อนุมัติโปรโมชั่น"],["Preconditions","เข้าสู่ระบบและมีสิทธิ์ในหน้าที่เกี่ยวข้อง"],["ขั้นตอนการทำงาน Step by step","1. เปิดรายการโปรโมชั่นหรือคำขออนุมัติ 2. ค้นหาและเลือกรายการ 3. ดูรายละเอียดและสถานะ 4. ดูผู้ส่ง/ผู้พิจารณา หมายเหตุ และเวลาเท่าที่ระบบบันทึก"],["Postconditions","ติดตามสถานะและประวัติการพิจารณาได้โดยไม่เปลี่ยนข้อมูล"],["Abnormal Paths","ไม่พบข้อมูลหรือโหลดไม่ได้ ระบบแสดงข้อความที่เหมาะสม"]])
set_table(d.tables[35], [["Actor","ประเภท","System Use Case ที่เกี่ยวข้อง"],["เจ้าหน้าที่ฝ่ายการตลาด","Person","U01, U02, U03, U05, U09"],["ผู้อนุมัติโปรโมชั่น","Person","U04, U09"],["ลูกค้าผู้ซื้อบัตร","Person","U06, U07, U08"]])
set_table(d.tables[36], [["Use Case ต้นทาง","ความสัมพันธ์","Use Case ปลายทาง","เหตุผล"],["U01 จัดการข้อมูล","include","U02 ตรวจข้อมูล","การบันทึกต้องผ่านการตรวจฟิลด์ รูปแบบ การอ้างอิง และรหัสไม่ซ้ำ"],["U01 จัดการข้อมูล","include","U03 สร้างคำขออนุมัติ","การสร้างหรือแก้ไขจะสร้างคำขอรออนุมัติทันที"],["U05 ลบโปรโมชั่น","extend","U01 จัดการข้อมูล","การลบเป็นการดำเนินการเพิ่มเติมกับรายการที่มีอยู่"],["U07 เลือกหรือกรอกรหัส","include","U06 ตรวจโปรโมชั่นที่ใช้ได้","ต้องตรวจเงื่อนไขก่อนนำโปรโมชั่นไปคำนวณ"],["U07 เลือกหรือกรอกรหัส","include","U08 คำนวณส่วนลด","เมื่อเลือกโปรโมชั่นแล้ว ระบบคำนวณส่วนลดสำหรับรายการจอง"]])

# ปรับข้อความที่อยู่ในย่อหน้านอกตาราง
replacements = {
    "U02 Customer Account Management":"U02 Employee Self Account and Password",
    "U03 Employee Account Management":"U03 Employee Account Management",
    "U04 Individual Employee Permission Management":"U04 Module Permission Management",
    "U05 Activity History Review":"U05 Activity and Password-reset Review",
    "U07 Apply the Highest-Discount Promotion":"U07 Select or Enter Promotion Code",
    "U08 Deduct or Restore Promotion Quota":"U08 Calculate Booking Discount",
    "U09 Review Promotion History and Results":"U09 Review Approval Status and History",
}
for p in d.paragraphs:
    for old,new in replacements.items():
        if old in p.text:
            for run in p.runs: run.text = run.text.replace(old,new)

d.save(TMP)

user_cases={
 "U01":("U01 Login / Logout",.42,.22,.34,.11),
 "U02":("U02 Self Account and Password",.58,.38,.38,.11),
 "U03":("U03 Manage Employee Accounts",.42,.55,.38,.11),
 "U04":("U04 Manage Module Permissions",.58,.70,.40,.11),
 "U05":("U05 Review Activity and Reset Requests",.42,.85,.42,.11),
}
img1=usecase_diagram((2692,2008),"User Management and Access Control System",[
 ("Employee",.08,.29,["U01","U02"]),("Administrator",.91,.57,["U01","U03","U04","U05"])
],user_cases,[("U03","U04","<<include>>")])

promo_cases={
 "U01":("U01 Manage Promotion",.38,.18,.34,.07),"U02":("U02 Validate Data",.62,.27,.31,.07),
 "U03":("U03 Create Pending Approval",.38,.36,.37,.07),"U04":("U04 Approve / Reject",.62,.46,.32,.07),
 "U05":("U05 Soft-delete Promotion",.38,.56,.35,.07),"U06":("U06 Check Eligible Promotions",.62,.65,.39,.07),
 "U07":("U07 Select or Enter Code",.38,.74,.36,.07),"U08":("U08 Calculate Booking Discount",.62,.83,.40,.07),
 "U09":("U09 Review Approval History",.38,.91,.38,.07),
}
img3=usecase_diagram((2516,2968),"Promotion Management System",[
 ("Marketing Staff",.07,.30,["U01","U05","U09"]),("Approver",.92,.44,["U04","U09"]),("Customer",.08,.69,["U06","U07","U08"])
],promo_cases,[("U01","U02","<<include>>"),("U01","U03","<<include>>"),("U07","U06","<<include>>"),("U07","U08","<<include>>")])

img2=flow_diagram((6568,12568),"User Management and Access Control - Activity Diagram",["Employee","System","Administrator"],{
 "s":(0,.09,"","start"),"login":(0,.16,"Enter credentials","task"),"check":(1,.24,"Validate account, password and status","task"),"ok":(1,.32,"Valid?","decision"),"menu":(1,.40,"Load module permissions and show allowed menu","task"),"self":(0,.49,"Manage own profile/password or request reset","task"),"admin":(2,.58,"Manage employee / permissions / reset requests","task"),"audit":(2,.68,"Review recorded activity","task"),"logout":(0,.79,"Logout","task"),"e":(1,.88,"","end")
},[("s","login",""),("login","check",""),("check","ok",""),("ok","menu","Yes"),("menu","self",""),("self","admin","Admin only"),("admin","audit",""),("audit","logout",""),("logout","e","")])
img4=flow_diagram((9208,9728),"Promotion Management - Activity Diagram",["Marketing Staff","System","Approver","Customer"],{
 "s":(0,.10,"","start"),"edit":(0,.17,"Create or edit supported promotion fields","task"),"validate":(1,.25,"Validate fields, dates, references and unique code","task"),"valid":(1,.33,"Valid?","decision"),"pending":(1,.41,"Save promotion and create pending approval","task"),"review":(2,.49,"Review and approve or reject with remark","task"),"approved":(1,.57,"Active and approved?","decision"),"eligible":(1,.65,"Filter by date, quota, concert, zone and minimum order","task"),"choose":(3,.73,"Select one promotion or enter code","task"),"discount":(1,.81,"Calculate discount and send discount_amount with booking","task"),"note":(1,.88,"Current version does not log promotion_id or update UsedQuota","task"),"e":(3,.94,"","end")
},[("s","edit",""),("edit","validate",""),("validate","valid",""),("valid","pending","Yes"),("pending","review",""),("review","approved",""),("approved","eligible","Yes"),("eligible","choose",""),("choose","discount",""),("discount","note",""),("note","e","")])

with zipfile.ZipFile(TMP,"r") as zin, zipfile.ZipFile(OUT,"w",zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data=zin.read(item.filename)
        if item.filename=="word/media/image1.png": data=img1
        elif item.filename=="word/media/image2.png": data=img2
        elif item.filename=="word/media/image3.png": data=img3
        elif item.filename=="word/media/image4.png": data=img4
        zout.writestr(item,data)
TMP.unlink(missing_ok=True)
print(OUT)

from pathlib import Path
from docx import Document
from pypdf import PdfReader
import zipfile

docx=Path(r"D:\SA\Mock-SA\เอกสาร SA jingjing ฉบับปรับให้ตรงระบบ.docx")
pdf=Path(r"D:\SA\Mock-SA\_rendered_revised.pdf")
d=Document(docx)
txt="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
stale=["จัดการบัญชีลูกค้า","ส่วนลดสูงสุดเพียงรายการเดียว","ไม่ต้องกรอกรหัสโปรโมชั่น","กลุ่มลูกค้า","ช่องทางขาย","ประเภทบัตร","ตัดหรือคืนจำนวนสิทธิ์","U06 ตรวจสิทธิ์โปรโมชั่น | include | U06"]
assert all(x not in txt for x in stale), [x for x in stale if x in txt]
required=["เลือกหรือกรอกรหัสโปรโมชั่น","soft delete","PromotionUsageLog","none, view หรือ edit","คำขอรีเซ็ตรหัสผ่าน"]
assert all(x in txt for x in required), [x for x in required if x not in txt]
with zipfile.ZipFile(docx) as z:
    media=[n for n in z.namelist() if n.startswith("word/media/")]
    assert len(media)==4
pages=len(PdfReader(pdf).pages)
assert pages==44, pages
assert len(list(Path(r"D:\SA\Mock-SA\_render_pages").glob("page-*.png")))==44
print(f"PASS docx={docx.name} tables={len(d.tables)} media={len(media)} pdf_pages={pages} stale_terms=0 required_terms={len(required)}")
